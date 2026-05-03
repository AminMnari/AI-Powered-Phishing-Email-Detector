from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> None:
    # Get project root (parent of src directory)
    PROJECT_ROOT = Path(__file__).parent.parent

    doc = Document()

    add_title(doc, "AI-Powered Phishing Email Detector")
    subtitle = doc.add_paragraph("Technical Report: Main Concepts, System Components, Functional Flow, and Existing Solutions")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line = doc.add_paragraph(f"Date: {date.today().isoformat()}")
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "1. Studied Topic: Main Concepts", level=1)
    doc.add_paragraph(
        "Phishing email detection is a cybersecurity classification problem where each email is categorized as malicious "
        "(phishing/spam) or legitimate (ham). The project uses supervised machine learning to learn patterns from labeled "
        "data and then predicts the risk level of new incoming emails."
    )

    add_heading(doc, "1.1 Core Cybersecurity Concepts", level=2)
    add_bullets(doc, [
        "Phishing: social-engineering attacks that trick users into revealing credentials or clicking malicious links.",
        "False Positive (FP): legitimate email predicted as phishing.",
        "False Negative (FN): phishing email predicted as legitimate.",
        "Precision: among predicted phishing emails, the proportion that is truly phishing.",
        "Recall: among true phishing emails, the proportion correctly detected.",
        "F1-score: harmonic mean of precision and recall.",
    ])

    add_heading(doc, "1.2 Core Machine Learning Concepts", level=2)
    add_bullets(doc, [
        "Text preprocessing: cleaning raw email content (HTML stripping, URL/email masking, normalization).",
        "Feature engineering: combining statistical text features (TF-IDF) with handcrafted phishing indicators.",
        "Model training: learning a decision boundary from labeled examples.",
        "Thresholding: converting phishing probability into a final class decision.",
        "Hybrid inference: combining ML predictions with rule-based logic to reduce false positives.",
    ])

    add_heading(doc, "2. Main Components of This Project", level=1)
    add_bullets(doc, [
        "Data layer: datasets in data/phishing_legitimate.csv and data/spam_assassin.csv.",
        "Preparation layer: schema normalization, label harmonization, and text cleaning.",
        "Feature layer: TF-IDF vectors + phishing/legitimacy custom signals.",
        "Model layer: Complement Naive Bayes and Logistic Regression (baseline and v2).",
        "Serving layer: Flask API + web interface for interactive testing.",
    ])

    add_heading(doc, "3. Functional Flow of the System", level=1)
    add_heading(doc, "3.1 Offline Training Flow", level=2)
    add_bullets(doc, [
        "Load and normalize both datasets.",
        "Clean text and generate cleaned_emails.csv.",
        "Build feature matrices (X) and labels (y).",
        "Split into train/test and train candidate models.",
        "Evaluate with accuracy, precision, recall, F1, and confusion matrices.",
        "Save vectorizer and best model artifacts for deployment.",
    ])

    add_heading(doc, "3.2 Online Prediction Flow", level=2)
    add_bullets(doc, [
        "User submits email text via UI or API.",
        "Server loads artifacts (lazy loading).",
        "Text is cleaned and transformed into TF-IDF + custom features.",
        "Model outputs phishing probability.",
        "Threshold and/or ensemble override is applied.",
        "JSON response returns label, confidence, and interpretation fields.",
    ])

    add_heading(doc, "4. Overview of Main Existing Solutions", level=1)
    add_bullets(doc, [
        "Rule-based email filters: fast but easy to evade.",
        "Classical ML classifiers: Naive Bayes, Logistic Regression, SVM, Random Forest.",
        "Deep learning / transformers: stronger semantic understanding, higher compute cost.",
        "SPF, DKIM, and DMARC: email authentication protocols for sender validation.",
        "Secure email gateways: commercial solutions integrating reputation and sandboxing.",
        "Hybrid systems: combine authentication, reputation, ML scoring, and heuristics.",
    ])

    add_heading(doc, "5. Project Positioning", level=1)
    doc.add_paragraph(
        "This project sits in the classical ML + hybrid inference category. Its main value is a complete pipeline from data "
        "preparation to deployment, with a practical focus on reducing false positives using ensemble logic."
    )

    output_path = PROJECT_ROOT / "AI_Powered_Phishing_Email_Detector_Report.docx"
    doc.save(output_path)
    print(f"Saved: {output_path.resolve()}")


if __name__ == "__main__":
    main()
